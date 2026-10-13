"""Regression tests for the findings of security audit run-1."""

import io
from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.config import Settings, get_settings
from app.models import BookingStatus, User, UserRole
from app.services import availability
from tests.conftest import DEFAULT_PASSWORD, auth_headers
from tests.factories import insert_booking, make_service, offer, tomorrow_at

API = "/api/v1"


@pytest.fixture
def world(session: Session, make_user: Callable[..., User]) -> dict[str, Any]:
    provider = make_user(role=UserRole.provider)
    service = make_service(session, duration=30)
    offer(session, provider, service)
    return {"provider": provider, "service": service, "admin": make_user(role=UserRole.admin)}


def _book(client: TestClient, world: dict[str, Any], customer: User, starts_at: datetime) -> Any:
    return client.post(
        f"{API}/bookings",
        json={
            "provider_id": str(world["provider"].id),
            "service_id": str(world["service"].id),
            "starts_at": starts_at.isoformat(),
        },
        headers=auth_headers(customer),
    )


# --- SECRET_KEY strength -------------------------------------------------------


@pytest.mark.parametrize("key", ["change-me-in-production-use-a-long-random-string", "short", ""])
def test_weak_or_placeholder_secret_key_is_rejected(key: str) -> None:
    with pytest.raises(ValidationError):
        Settings(secret_key=key, database_url="postgresql://x/y", _env_file=None)  # type: ignore[call-arg]


def test_strong_secret_key_is_accepted() -> None:
    settings = Settings(secret_key="k" * 48, database_url="postgresql://x/y", _env_file=None)  # type: ignore[call-arg]
    assert settings.secret_key == "k" * 48


# --- Booking requires a verified email ----------------------------------------


def test_unverified_customer_cannot_book(
    client: TestClient, world: dict[str, Any], make_user: Callable[..., User]
) -> None:
    unverified = make_user(verified=False)
    response = _book(client, world, unverified, tomorrow_at(10))
    assert response.status_code == 403
    assert response.json() == {"detail": "email_not_verified"}


# --- Per-customer cap on active bookings ---------------------------------------


def test_customer_active_booking_cap(
    client: TestClient,
    world: dict[str, Any],
    make_user: Callable[..., User],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "max_active_bookings_per_customer", 2)
    customer = make_user()
    assert _book(client, world, customer, tomorrow_at(9)).status_code == 201
    assert _book(client, world, customer, tomorrow_at(10)).status_code == 201
    blocked = _book(client, world, customer, tomorrow_at(11))
    assert blocked.status_code == 409
    assert blocked.json() == {"detail": "too_many_active_bookings"}
    # Another customer can still book: the provider was not hoarded.
    assert _book(client, world, make_user(), tomorrow_at(11)).status_code == 201


def test_cancelled_and_past_bookings_do_not_count(
    client: TestClient,
    session: Session,
    world: dict[str, Any],
    make_user: Callable[..., User],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(get_settings(), "max_active_bookings_per_customer", 1)
    customer = make_user()
    provider, service = world["provider"], world["service"]
    insert_booking(session, customer, provider, service, tomorrow_at(8), BookingStatus.cancelled)
    insert_booking(
        session,
        customer,
        provider,
        service,
        datetime.now(UTC) - timedelta(days=1),
        BookingStatus.confirmed,
    )
    assert _book(client, world, customer, tomorrow_at(10)).status_code == 201


# --- full_name control characters and the XLSX report ---------------------------


@pytest.mark.parametrize("name", ["Cliente\x01Prueba", "X\x0bY", "Tab\tName", "Del\x7f"])
def test_full_name_rejects_control_characters(client: TestClient, name: str) -> None:
    response = client.post(
        f"{API}/auth/register",
        json={"email": "ctrl@example.com", "password": DEFAULT_PASSWORD, "full_name": name},
    )
    assert response.status_code == 422


def test_users_me_rejects_control_characters(
    client: TestClient, make_user: Callable[..., User]
) -> None:
    user = make_user()
    response = client.patch(
        f"{API}/users/me", json={"full_name": "Ver\x1fCinco"}, headers=auth_headers(user)
    )
    assert response.status_code == 422


def test_xlsx_report_survives_stored_control_characters(
    client: TestClient,
    session: Session,
    world: dict[str, Any],
    make_user: Callable[..., User],
) -> None:
    """Rows written before the schema check existed must not break the export."""
    customer = make_user()
    customer.full_name = "Legacy\x1fName"
    insert_booking(session, customer, world["provider"], world["service"], tomorrow_at(10))
    day = tomorrow_at(10).date()
    response = client.get(
        f"{API}/admin/reports",
        params={"format": "xlsx", "date_from": str(day), "date_to": str(day)},
        headers=auth_headers(world["admin"]),
    )
    assert response.status_code == 200
    detail = load_workbook(io.BytesIO(response.content))["Reservas"]
    assert detail.cell(2, 3).value == "LegacyName"


# --- Provider availability caps -------------------------------------------------


def test_availability_rules_are_capped(
    client: TestClient, world: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(availability, "MAX_RULES", 2)
    rules = [{"weekday": d, "start_time": "08:00:00", "end_time": "09:00:00"} for d in range(3)]
    response = client.put(
        f"{API}/providers/me/availability", json=rules, headers=auth_headers(world["provider"])
    )
    assert response.status_code == 422
    assert response.json() == {"detail": "too_many_rules"}


def test_future_time_off_is_capped(
    client: TestClient, world: dict[str, Any], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(availability, "MAX_FUTURE_TIME_OFF", 2)
    headers = auth_headers(world["provider"])

    def add(day: int) -> Any:
        start = datetime.now(UTC) + timedelta(days=day)
        return client.post(
            f"{API}/providers/me/time-off",
            json={
                "starts_at": start.isoformat(),
                "ends_at": (start + timedelta(hours=1)).isoformat(),
            },
            headers=headers,
        )

    assert add(3).status_code == 201
    assert add(4).status_code == 201
    blocked = add(5)
    assert blocked.status_code == 409
    assert blocked.json() == {"detail": "too_many_time_off"}
