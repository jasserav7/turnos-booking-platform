import io
from collections.abc import Callable
from datetime import UTC, datetime, timedelta
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import BookingStatus, User, UserRole
from tests.conftest import auth_headers
from tests.factories import insert_booking, make_service, offer, tomorrow_at

URL = "/api/v1/admin/reports"


@pytest.fixture
def world(session: Session, make_user: Callable[..., User]) -> dict[str, Any]:
    provider = make_user(role=UserRole.provider)
    provider.full_name = "Laura Gómez"
    customer = make_user()
    customer.full_name = '=HYPERLINK("http://evil","x")'
    service = make_service(session, price_cents=4_500_000)
    offer(session, provider, service)
    admin = make_user(role=UserRole.admin)
    insert_booking(session, customer, provider, service, tomorrow_at(10), BookingStatus.completed)
    insert_booking(session, customer, provider, service, tomorrow_at(12))
    return {"admin": admin, "customer": customer, "provider": provider}


def _range() -> dict[str, str]:
    day = tomorrow_at(10).date()
    return {"date_from": str(day - timedelta(days=1)), "date_to": str(day + timedelta(days=1))}


def _get(client: TestClient, world: dict[str, Any], **params: str) -> Any:
    return client.get(URL, params={**_range(), **params}, headers=auth_headers(world["admin"]))


def test_csv_report(client: TestClient, world: dict[str, Any]) -> None:
    response = _get(client, world, format="csv")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert 'attachment; filename="reporte-reservas_' in response.headers["content-disposition"]
    assert response.content.startswith(b"\xef\xbb\xbf")
    lines = response.content.decode("utf-8-sig").strip().split("\r\n")
    assert lines[0].split(";")[:3] == ["Fecha", "Hora", "Cliente"]
    assert len(lines) == 3
    assert "Laura Gómez" in lines[1]
    assert "Completada" in lines[1] and "Pendiente" in lines[2]


def test_csv_neutralizes_formulas(client: TestClient, world: dict[str, Any]) -> None:
    text = _get(client, world, format="csv").content.decode("utf-8-sig")
    assert ";=HYPERLINK" not in text
    assert "'=HYPERLINK" in text


def test_xlsx_report(client: TestClient, world: dict[str, Any]) -> None:
    response = _get(client, world, format="xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    wb = load_workbook(io.BytesIO(response.content))
    assert wb.sheetnames == ["Resumen", "Reservas"]
    detail = wb["Reservas"]
    assert detail.max_row == 3
    customer_cell = detail.cell(2, 3)
    assert customer_cell.data_type == "s"
    assert str(customer_cell.value).startswith("'=")
    summary = {row[0].value: row[1].value for row in wb["Resumen"].iter_rows(min_row=3)}
    assert summary["Reservas en el período"] == 2
    assert summary["Ingresos estimados (completadas)"] == 45_000


def test_pdf_report(client: TestClient, world: dict[str, Any]) -> None:
    response = _get(client, world, format="pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


def test_empty_range_still_renders(client: TestClient, world: dict[str, Any]) -> None:
    far = (datetime.now(UTC) + timedelta(days=200)).date()
    params = {"date_from": str(far), "date_to": str(far)}
    for fmt in ("csv", "xlsx", "pdf"):
        response = client.get(
            URL, params={**params, "format": fmt}, headers=auth_headers(world["admin"])
        )
        assert response.status_code == 200, fmt


def test_defaults_to_csv_last_30_days(client: TestClient, world: dict[str, Any]) -> None:
    response = client.get(URL, headers=auth_headers(world["admin"]))
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")


def test_invalid_ranges(client: TestClient, world: dict[str, Any]) -> None:
    headers = auth_headers(world["admin"])
    reversed_range = client.get(
        URL, params={"date_from": "2026-02-01", "date_to": "2026-01-01"}, headers=headers
    )
    assert reversed_range.status_code == 422
    assert reversed_range.json() == {"detail": "invalid_range"}
    too_large = client.get(
        URL, params={"date_from": "2024-01-01", "date_to": "2026-01-01"}, headers=headers
    )
    assert too_large.status_code == 422
    assert too_large.json() == {"detail": "range_too_large"}
    assert client.get(URL, params={"format": "docx"}, headers=headers).status_code == 422


def test_only_admins_download_reports(client: TestClient, world: dict[str, Any]) -> None:
    for actor in (world["customer"], world["provider"]):
        assert client.get(URL, headers=auth_headers(actor)).status_code == 403
    assert client.get(URL).status_code == 401
