import csv
import io
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.orm import Session, aliased

from app.core.config import get_settings
from app.models import Booking, BookingStatus, Service, User, UserRole
from app.services.permissions import ensure_roles

DEFAULT_DAYS = 30
MAX_RANGE_DAYS = 366
MAX_ROWS = 10_000

STATUS_LABELS = {
    BookingStatus.pending: "Pendiente",
    BookingStatus.confirmed: "Confirmada",
    BookingStatus.cancelled: "Cancelada",
    BookingStatus.completed: "Completada",
}

COLUMNS = [
    "Fecha",
    "Hora",
    "Cliente",
    "Correo del cliente",
    "Profesional",
    "Servicio",
    "Duración (min)",
    "Precio (COP)",
    "Estado",
]
PRICE_COLUMN = 8  # 1-based index of "Precio (COP)"
ACCENT_HEX = "8C2A87"


@dataclass(frozen=True)
class ReportRow:
    day: date
    start: time
    customer_name: str
    customer_email: str
    provider_name: str
    service_name: str
    duration_minutes: int
    price: int
    status: BookingStatus


@dataclass(frozen=True)
class Report:
    date_from: date
    date_to: date
    rows: list[ReportRow]

    @property
    def by_status(self) -> dict[BookingStatus, int]:
        counts = {s: 0 for s in BookingStatus}
        for row in self.rows:
            counts[row.status] += 1
        return counts

    @property
    def revenue(self) -> int:
        return sum(r.price for r in self.rows if r.status == BookingStatus.completed)


def _bad_range(detail: str) -> HTTPException:
    return HTTPException(status_code=422, detail=detail)


def collect_report(
    db: Session, actor: User, date_from: date | None, date_to: date | None
) -> Report:
    ensure_roles(actor, UserRole.admin)
    tz = ZoneInfo(get_settings().app_timezone)
    date_to = date_to or datetime.now(tz).date()
    date_from = date_from or date_to - timedelta(days=DEFAULT_DAYS - 1)
    if date_from > date_to:
        raise _bad_range("invalid_range")
    if (date_to - date_from).days >= MAX_RANGE_DAYS:
        raise _bad_range("range_too_large")

    window_start = datetime.combine(date_from, time.min, tzinfo=tz).astimezone(UTC)
    window_end = datetime.combine(date_to + timedelta(days=1), time.min, tzinfo=tz).astimezone(UTC)
    customer, provider = aliased(User), aliased(User)
    query = (
        select(Booking, customer, provider, Service)
        .join(customer, customer.id == Booking.customer_id)
        .join(provider, provider.id == Booking.provider_id)
        .join(Service, Service.id == Booking.service_id)
        .where(Booking.starts_at >= window_start, Booking.starts_at < window_end)
        .order_by(Booking.starts_at, Booking.id)
        .limit(MAX_ROWS)
    )
    rows = []
    for booking, cust, prov, service in db.execute(query):
        local = booking.starts_at.astimezone(tz)
        rows.append(
            ReportRow(
                day=local.date(),
                start=local.time().replace(second=0, microsecond=0),
                customer_name=cust.full_name,
                customer_email=cust.email,
                provider_name=prov.full_name,
                service_name=service.name,
                duration_minutes=service.duration_minutes,
                price=service.price_cents // 100,
                status=booking.status,
            )
        )
    return Report(date_from=date_from, date_to=date_to, rows=rows)


def _safe(value: str) -> str:
    """Neutralize spreadsheet formulas and drop characters XLSX cannot store.

    User-controlled text must never be evaluated, nor make the export fail.
    """
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return f"'{value}" if value[:1] in ("=", "+", "-", "@", "\t", "\r") else value


def _cells(row: ReportRow) -> list[str | int]:
    return [
        row.day.isoformat(),
        row.start.strftime("%H:%M"),
        _safe(row.customer_name),
        _safe(row.customer_email),
        _safe(row.provider_name),
        _safe(row.service_name),
        row.duration_minutes,
        row.price,
        STATUS_LABELS[row.status],
    ]


def _money(value: int) -> str:
    return "$ " + f"{value:,}".replace(",", ".")


def _summary(report: Report) -> list[tuple[str, int, bool]]:
    """(label, value, is_money)"""
    return [
        ("Reservas en el período", len(report.rows), False),
        *((STATUS_LABELS[s], n, False) for s, n in report.by_status.items()),
        ("Ingresos estimados (completadas)", report.revenue, True),
    ]


def render_csv(report: Report) -> bytes:
    """UTF-8 with BOM and ';' so Excel in Spanish locales opens accents and columns correctly."""
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", lineterminator="\r\n")
    writer.writerow(COLUMNS)
    for row in report.rows:
        writer.writerow(_cells(row))
    return out.getvalue().encode("utf-8-sig")


def render_xlsx(report: Report) -> bytes:
    wb = Workbook()
    detail = wb.active
    assert detail is not None
    detail.title = "Reservas"
    detail.append(COLUMNS)
    for row in report.rows:
        detail.append(_cells(row))
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=ACCENT_HEX)
        cell.alignment = Alignment(vertical="center")
    for r in range(2, detail.max_row + 1):
        detail.cell(r, PRICE_COLUMN).number_format = '"$" #,##0'
    for index, title in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        longest = max(len(title), *(len(str(c.value or "")) for c in detail[letter]))
        detail.column_dimensions[letter].width = min(longest + 3, 40)
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    summary = wb.create_sheet("Resumen", 0)
    summary.append(["Reporte de reservas", f"{report.date_from} a {report.date_to}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    for label, value, is_money in _summary(report):
        summary.append([label, value])
        if is_money:
            summary.cell(summary.max_row, 2).number_format = '"$" #,##0'
    summary.column_dimensions["A"].width = 42
    summary.column_dimensions["B"].width = 24

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(report: Report) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        title="Reporte de reservas",
    )
    styles = getSampleStyleSheet()
    accent = colors.HexColor(f"#{ACCENT_HEX}")
    story: list[Any] = [
        Paragraph("Reporte de reservas", styles["Title"]),
        Paragraph(f"Período: {report.date_from} a {report.date_to}", styles["Normal"]),
        Spacer(1, 0.4 * cm),
    ]
    summary_table = Table(
        [[label, _money(v) if money else str(v)] for label, v, money in _summary(report)],
        colWidths=[8 * cm, 4 * cm],
        hAlign="LEFT",
    )
    summary_table.setStyle(
        TableStyle(
            [
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ]
        )
    )
    story += [summary_table, Spacer(1, 0.6 * cm)]

    if not report.rows:
        story.append(Paragraph("No hay reservas en este período.", styles["Normal"]))
    else:
        head = [c.replace(" (COP)", "") for c in COLUMNS]
        data = [head] + [
            [
                r.day.isoformat(),
                r.start.strftime("%H:%M"),
                r.customer_name,
                r.customer_email,
                r.provider_name,
                r.service_name,
                str(r.duration_minutes),
                _money(r.price),
                STATUS_LABELS[r.status],
            ]
            for r in report.rows
        ]
        widths_cm = [2.2, 1.4, 3.6, 5.2, 3.6, 3.6, 2.0, 2.4, 2.2]
        table = Table(data, repeatRows=1, colWidths=[w * cm for w in widths_cm])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), accent),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F8F4F9")],
                    ),
                    ("ALIGN", (6, 0), (7, -1), "RIGHT"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        story.append(table)
    doc.build(story)
    return out.getvalue()
